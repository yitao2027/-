"""
勺子Claw 邀请码生成器
- 11位复杂邀请码（数字+字母 校验码）
- 数量：500
- 输出：xlsx 销售表 + ts 校验列表
"""
import secrets
import string
import json
import hashlib
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 字符集：去除易混淆字符 0/O/1/I/l
CHARSET = "23456789ABCDEFGHJKLMNPQRSTUVWXYZ"  # 32个字符

def gen_one_code() -> str:
    """生成 11 位邀请码：8位随机 + 3位校验"""
    # 8位主体随机
    body = ''.join(secrets.choice(CHARSET) for _ in range(8))
    # 3位 sha256 截取的伪校验码（仍走字符集，更像内部码）
    h = hashlib.sha256(body.encode()).hexdigest().upper()
    check = ''
    for ch in h:
        if ch in CHARSET:
            check += ch
            if len(check) >= 3:
                break
    if len(check) < 3:
        check = (check + 'X' * 3)[:3]
    return body + check  # 共 11 位

def main():
    out_dir = Path(r"C:/Users/Administrator/WorkBuddy/20260513150254/shaoziclaw-app/deliverables/v5.5.20")
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1. 生成 500 个唯一码
    codes = set()
    while len(codes) < 500:
        codes.add(gen_one_code())
    codes = sorted(codes)

    # 2. 写 xlsx — 销售用表
    wb = Workbook()
    ws = wb.active
    ws.title = "勺子Claw内测邀请码"

    headers = ["序号", "邀请码", "邀请人姓名", "邀请人手机号", "邀请人企业名称", "邀请人岗位", "邀请人备注", "状态", "发放时间"]
    ws.append(headers)

    # 表头样式
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E7D5F")
    thin = Side(border_style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # 数据行
    code_font = Font(name="Consolas", size=10)
    for i, code in enumerate(codes, 1):
        row = [i, code, "", "", "", "", "", "未发放", ""]
        ws.append(row)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=i + 1, column=col_idx)
            c.border = border
            c.alignment = center
            if col_idx == 2:
                c.font = code_font
                c.fill = PatternFill("solid", fgColor="F4F9F6")

    # 列宽
    widths = [6, 16, 14, 16, 22, 14, 24, 10, 18]
    for col_idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # 冻结首行
    ws.freeze_panes = "A2"

    # 增加说明 sheet
    ws2 = wb.create_sheet("使用说明")
    instructions = [
        ["勺子Claw 邀请码使用说明"],
        [],
        ["1. 总数: 500 个 11 位邀请码（数字+大写字母）"],
        ["2. 字符集: 已去除易混淆字符 0/O/1/I/l"],
        ["3. 每个码唯一，校验码由系统 sha256 生成，可后端校验"],
        ["4. 发放流程: 销售/同事 → 填写邀请人信息 → 改状态为 已发放"],
        ["5. 用户使用: 登录页选择「邀请码登录」 → 粘贴 11 位码"],
        ["6. 失效条件: 单码绑定一个手机号后失效（后端逻辑）"],
        [],
        ["生成时间: 2026-05-14"],
        ["项目版本: 勺子Claw v5.5.20"],
    ]
    for row in instructions:
        ws2.append(row)
    ws2.column_dimensions["A"].width = 80
    ws2["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="2E7D5F")

    xlsx_path = out_dir / "勺子Claw_内测邀请码_500个_v5.5.20.xlsx"
    wb.save(xlsx_path)
    print(f"[OK] xlsx 已生成: {xlsx_path}")

    # 3. 写 ts 校验列表（前端用）
    ts_path = out_dir / "INVITE_CODES.ts"
    ts_content = "// 自动生成 - 勺子Claw v5.5.20 - 500个内测邀请码\n"
    ts_content += "// 生成时间: 2026-05-14\n"
    ts_content += "// 字符集: 32字符（去除0/O/1/I/l），11位（8位主体+3位校验）\n\n"
    ts_content += "export const VALID_INVITE_CODES: string[] = [\n"
    for code in codes:
        ts_content += f'  "{code}",\n'
    ts_content += "];\n\n"
    ts_content += "export const INVITE_CODE_COUNT = " + str(len(codes)) + ";\n"
    ts_content += "export const INVITE_CODE_PATTERN = /^[23456789A-HJ-NP-Z]{11}$/;\n"
    ts_path.write_text(ts_content, encoding="utf-8")
    print(f"[OK] ts 校验列表: {ts_path}")

    # 4. JSON备份
    json_path = out_dir / "invite_codes.json"
    json_path.write_text(json.dumps({"version": "v5.5.20", "count": len(codes), "codes": codes}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OK] JSON 备份: {json_path}")

    print(f"\n样例（前5个）: {codes[:5]}")
    print(f"样例（最后5个）: {codes[-5:]}")

if __name__ == "__main__":
    main()
